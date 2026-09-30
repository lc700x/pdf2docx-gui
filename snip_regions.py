"""Replace converted tables and display equations with crops of the PDF.

pdf2docx reads a PDF as positioned glyphs, so it has no notion of a formula: a
displayed equation comes out as loose fragments whose summation limits and
subscripts drift, and a table built from those fragments drifts with them. The
page itself is already correct, so for these regions the faithful thing is to
show the page.

Body text is left exactly as converted, and so are captions, which sit outside
the region and stay searchable.

Regions are found on the PDF:

* **tables** by their rules — a booktabs table has no grid for a table finder
  to latch onto, but its horizontal rules bound it precisely;
* **display equations** by being set apart — indented, short of the right
  margin, and mostly set in the math fonts.

Each region is then matched to the converted content by text, so a crop
replaces the elements it actually came from.
"""

import re
from collections import Counter
from difflib import SequenceMatcher

import fitz
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

# The math and symbol fonts a TeX document sets formulas in. A run of body text
# never reaches far into these, so the share of characters drawn with them
# separates a formula from a sentence that merely mentions a variable.
MATH_FONT = re.compile(
    r'(NewTX|txsy|txmi|txex|txsym|CM(MI|SY|EX)|MathJax|Math[A-Z]|'
    r'STIX|Euler|msam|msbm|rsfs|wasy)', re.I)

RULE_MAX_HEIGHT = 2.5      # a rule is a filled box this thin
RULE_MIN_WIDTH = 40        # ... and at least this wide
EQ_MIN_MATH = 0.25         # share of characters set in a math font
EQ_MIN_INDENT = 24         # how far in from the text margin a display line sits
EQ_MERGE_GAP = 12         # vertical gap that still belongs to one formula
CAPTION = re.compile(
    r'^\s*('
    r'(Table|Figure|Fig\.?)\s*\d+\s*[:.]'    # "Table 3:", "Figure 2."
    r'|\([a-z]\)\s'                          # "(a) ", the caption of a panel
    r')', re.I)


def _base_font(name):
    return name.split('+')[-1]


def _blocks(page):
    """Text blocks with the geometry and the math share this module needs."""
    out = []
    for b in page.get_text('dict')['blocks']:
        if b['type'] != 0:
            continue
        text = ''
        math = 0
        for line in b['lines']:
            for span in line['spans']:
                text += span['text']
                if MATH_FONT.search(_base_font(span['font'])):
                    math += len(span['text'])
        body = len(text.strip())
        out.append({
            'rect': fitz.Rect(b['bbox']),
            'text': text,
            'math': math / body if body else 0.0,
        })
    return out


def _text_frame(blocks, page):
    """Left and right edge of the body text, read off the page itself."""
    wide = [b['rect'] for b in blocks if b['rect'].width > page.rect.width * 0.5]
    if not wide:
        return page.rect.x0 + 72, page.rect.x1 - 72
    return min(r.x0 for r in wide), max(r.x1 for r in wide)


def _rules(page):
    out = []
    for drawing in page.get_drawings():
        rect = fitz.Rect(drawing['rect'])
        if rect.height <= RULE_MAX_HEIGHT and rect.width >= RULE_MIN_WIDTH:
            out.append(rect)
    return sorted(out, key=lambda r: r.y0)


def _caption_between(blocks, top, bottom):
    return any(CAPTION.match(b['text'])
               for b in blocks if top <= b['rect'].y0 and b['rect'].y1 <= bottom)


def _same_table_width(a, b):
    """Whether two rules could frame or divide the same table.

    Every rule of one table spans its full measure or, for a rule under a group
    of columns, a part of it — so one range contains the other. Two tables set
    to different measures, such as the panels of a confusion matrix, do not.
    """
    overlap = min(a.x1, b.x1) - max(a.x0, b.x0)
    if overlap <= 0:
        return False
    return overlap >= min(a.width, b.width) * 0.9


def find_tables(page, blocks, left, right):
    """Table regions, bounded by the rules that frame each one."""
    rules = _rules(page)
    if not rules:
        return []
    groups = [[rules[0]]]
    for rule in rules[1:]:
        group = groups[-1]
        widest = max(group, key=lambda r: r.width)
        gap_top, gap_bottom = group[-1].y1, rule.y0
        if (_caption_between(blocks, gap_top, gap_bottom)
                or not _same_table_width(widest, rule)):
            groups.append([rule])
        else:
            group.append(rule)

    regions = []
    for group in groups:
        top = min(r.y0 for r in group)
        bottom = max(r.y1 for r in group)
        box = fitz.Rect(min(r.x0 for r in group), top,
                        max(r.x1 for r in group), bottom)
        content = False
        for b in blocks:                      # take in the cells between rules
            r = b['rect']
            if top - 2 <= r.y0 and r.y1 <= bottom + 2 and b['text'].strip():
                box |= r
                content = True
        if content and box.height >= 6:       # a lone rule frames no table
            regions.append(box)
    return regions


def find_equations(page, blocks, left, right, tables):
    """Display equations: set apart from the margin and mostly math."""
    def set_apart(r):
        return (r.x0 >= left + EQ_MIN_INDENT   # not starting at the margin
                and r.x1 <= right - 2          # not running to the margin
                and not any(r.intersects(t) for t in tables))

    candidates = [b for b in blocks if b['text'].strip() and set_apart(b['rect'])]
    seeds = [b['rect'] for b in candidates if b['math'] >= EQ_MIN_MATH]
    if not seeds:
        return []

    # A formula is laid out in pieces, and a piece can be almost free of math
    # characters on its own — "E_total = 10 log" beside the summation it sums.
    # Any set-apart piece level with a seed belongs to the same formula.
    regions = [fitz.Rect(r) for r in seeds]
    taken = [False] * len(candidates)
    for _ in range(len(candidates)):
        grew = False
        for i, b in enumerate(candidates):
            if taken[i]:
                continue
            r = b['rect']
            for k, region in enumerate(regions):
                if r.y1 > region.y0 - EQ_MERGE_GAP and r.y0 < region.y1 + EQ_MERGE_GAP:
                    regions[k] = region | r   # Rect |= rebinds, it does not widen
                    taken[i] = True
                    grew = True
                    break
        if not grew:
            break

    regions.sort(key=lambda r: r.y0)
    merged = []
    for r in regions:
        if merged and r.y0 <= merged[-1].y1 + EQ_MERGE_GAP:
            merged[-1] |= r
        else:
            merged.append(fitz.Rect(r))
    return [r for r in merged if r.height >= 6]


def find_regions(page):
    """Every region on this page that should be shown as a crop, top to bottom."""
    blocks = _blocks(page)
    left, right = _text_frame(blocks, page)
    tables = find_tables(page, blocks, left, right)
    equations = find_equations(page, blocks, left, right, tables)
    regions = [('table', r) for r in tables] + [('equation', r) for r in equations]
    regions.sort(key=lambda item: item[1].y0)
    return regions, (left, right)


# --------------------------------------------------------------------------
# matching a region to the elements pdf2docx produced from it
# --------------------------------------------------------------------------

def _key(text):
    """Comparison key: the characters, with the spacing thrown away.

    Spacing is exactly what the conversion gets wrong, so it cannot be part of
    the comparison; both sides come from the same extractor, so the characters
    themselves line up.
    """
    return re.sub(r'\s+', '', text)


def _element_text(element):
    return ''.join(node.text or '' for node in element.iter(qn('w:t')))


def _pages(body, count):
    """The converted content grouped by the PDF page it came from.

    Section breaks look like the obvious divider but do not line up: pdf2docx
    opens extra sections around floats, so several sections belong to no page
    of their own. The printed folio is exact — each page ends with its own
    number — so the groups are cut there instead.
    """
    groups = [[] for _ in range(count)]
    current = []
    page = 0
    for child in body.iterchildren():
        if child.tag not in (qn('w:p'), qn('w:tbl')):
            continue
        current.append(child)
        if (page < count and child.tag == qn('w:p')
                and _element_text(child).strip() == str(page + 1)):
            groups[page] = current
            current = []
            page += 1
    if current and page < count:
        groups[page] = current
    return groups


MIN_ANCHOR = 6      # characters an element needs before it can bound a run
MIN_COVER = 0.75     # how completely one text must account for the other
MIN_COVER_TABLE = 0.5  # ... but a converted table row is never a sentence
INTERIOR_COVER = 0.6   # ... and the bar for an element the region merely brackets


SHORT_TARGET = 60   # characters: below this a region is a formula, not prose


def _cover(key, target):
    """How completely one of two texts accounts for the other.

    Containment is too strict in both directions: a table continued from the
    previous page is one element holding more than its region shows, and a
    caption fragment holds less. So the test is that one of the two is nearly
    all inside the other — not merely that they overlap. Overlap alone marks
    any paragraph that mentions the same method as part of the table beside
    it, and deletes it.

    A formula is the exception: it is short, and its pieces are not laid down
    in reading order, so for a short region the characters are compared
    without regard to order — against the longer of the two, so that a
    paragraph of prose cannot pass by containing a handful of symbols.
    """
    if len(key) < MIN_ANCHOR or not target:
        return 0.0
    if key in target:
        return 1.0
    matcher = SequenceMatcher(None, key, target, autojunk=False)
    common = sum(block.size for block in matcher.get_matching_blocks())
    score = max(common / len(key), common / len(target))
    if len(target) <= SHORT_TARGET:
        shared = sum((Counter(key) & Counter(target)).values())
        score = max(score, shared / max(len(key), len(target)))
    return score


def _word_cover(text, target):
    """Compare prose by words so a fragment cannot match inside a word."""
    token_re = r"[^\W\d_]+(?:['-][^\W\d_]+)*"
    source = [word.casefold() for word in re.findall(token_re, text)]
    wanted = [word.casefold() for word in re.findall(token_re, target)]
    if not source or not wanted:
        return _cover(_key(text), _key(target))
    matcher = SequenceMatcher(None, source, wanted, autojunk=False)
    common = sum(block.size for block in matcher.get_matching_blocks())
    return max(common / len(source), common / len(wanted))


def _element_cover(element, target):
    text = _element_text(element)
    if element.tag == qn('w:p'):
        return _word_cover(text, target)
    return _cover(_key(text), _key(target))


def _assign(elements, regions):
    """Which elements each region was converted into.

    Every element goes to the region it matches best, never to the first that
    will have it: neighbouring formulas share most of their symbols, so a
    region asked on its own would happily swallow the one after it and take
    that content out of the document.

    A caption is never taken. It describes the region from outside it and
    repeats much of its wording, which is exactly what makes it score well, so
    it has to be held out by name or it disappears into the crop.

    Returns, per region, the list of element indices to replace, or None.
    """
    captions = [bool(CAPTION.match(_element_text(e).strip())) for e in elements]
    # A converted table is a table; prose never is. That alone separates a row
    # of a table from a sentence that happens to name the same method, which no
    # amount of text similarity does — the two overlap by about as much.
    bars = [MIN_COVER_TABLE if e.tag == qn('w:tbl') else MIN_COVER
            for e in elements]

    owner = [None] * len(elements)
    for i, element in enumerate(elements):
        if captions[i]:
            continue
        best, score = None, bars[i]
        for r, target in enumerate(regions):
            cover = _element_cover(element, target)
            if cover > score:
                best, score = r, cover
        owner[i] = best

    groups = []
    for r in range(len(regions)):
        mine = [i for i, o in enumerate(owner) if o == r]
        if not mine:
            groups.append(None)
            continue
        # Between the matched elements sit two different things: fragments of
        # the same region — a wrapped cell, a stray line — and paragraphs of
        # body text that the region's own pieces happen to bracket. Only the
        # first kind may go, so an unclaimed element has to look like the
        # region's own text, or be too short to tell, before it is removed.
        keep = []
        for i in range(mine[0], mine[-1] + 1):
            if captions[i] or owner[i] not in (r, None):
                continue
            if owner[i] == r:
                keep.append(i)
                continue
            key = _key(_element_text(elements[i]))
            bar = MIN_COVER_TABLE if bars[i] == MIN_COVER_TABLE else INTERIOR_COVER
            if (len(key) < MIN_ANCHOR
                    or _element_cover(elements[i], regions[r]) >= bar):
                keep.append(i)
        groups.append(keep)
    return groups





def _remove_cropped_table_rows(table, target):
    """Drop rows covered by the crop but retain adjacent text and headings."""
    target = _key(target).casefold()
    removed = 0
    for row in list(table.findall(qn('w:tr'))):
        cells = [_key(_element_text(cell)).casefold()
                 for cell in row.findall(qn('w:tc'))]
        cells = [key for key in cells
                 if len(key) >= 3 or (len(key) >= 2 and key.isdigit())]
        if not cells:
            continue
        matched = sum(key in target or _cover(key, target) >= 0.9
                      for key in cells)
        if matched * 3 >= len(cells) * 2:
            table.remove(row)
            removed += 1
    return removed


def _insert_image(elements, indices, png, width_pt, document, target,
                  description='', preserve_table_rows=False):
    """Put the crop where the matched elements were, and drop them."""
    anchor = elements[indices[0]]
    new_p = anchor.makeelement(qn('w:p'), {})
    anchor.addprevious(new_p)

    from docx.text.paragraph import Paragraph
    paragraph = Paragraph(new_p, document)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = paragraph.add_run().add_picture(png, width=Pt(width_pt))
    if description:
        picture._inline.docPr.set('descr', description)

    for index in indices:
        element = elements[index]
        parent = element.getparent()
        if (parent is not None and element.tag == qn('w:tbl')
                and preserve_table_rows):
            _remove_cropped_table_rows(element, target)
            if not element.findall(qn('w:tr')):
                parent.remove(element)
        elif parent is not None:
            parent.remove(element)
    return new_p


def snip(pdf_path, docx_path, dpi=300, tables=True, equations=True,
         progress=None):
    """Swap each detected region in `docx_path` for a crop of `pdf_path`.

    Returns (replaced, missed): how many regions were swapped in, and how many
    could not be matched to converted content and so were left alone.
    """
    from docx import Document

    pdf = fitz.open(pdf_path)
    document = Document(docx_path)
    body = document.element.body
    parts = _pages(body, pdf.page_count)

    zoom = dpi / 72.0
    replaced = missed = 0
    tmp = []

    try:
        for index, page in enumerate(pdf):
            if index >= len(parts):
                break
            regions, (left, right) = find_regions(page)
            regions = [(kind, rect) for kind, rect in regions
                       if (kind == 'table' and tables)
                       or (kind == 'equation' and equations)]
            if not regions:
                continue

            elements = parts[index]
            wanted = [page.get_textbox(rect + (-1, -1, 1, 1))
                      for _, rect in regions]
            groups = _assign(elements, wanted)

            # Later regions first, so that removing one region's elements does
            # not disturb the elements of the regions still to come.
            for (kind, rect), indices, target in sorted(
                    zip(regions, groups, wanted),
                    key=lambda item: -(item[1][0] if item[1] else 0)):
                if not indices:
                    missed += 1
                    continue

                pad = 3 if kind == 'equation' else 5
                box = (rect + (-pad, -pad, pad, pad)) & page.rect
                pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom),
                                         clip=box)
                png = f'{docx_path}.snip-{index:04d}-{indices[0]:04d}.png'
                pixmap.save(png)
                tmp.append(png)

                width = min(box.width, right - left)
                description = page.get_textbox(box).strip()
                _insert_image(elements, indices, png, width, document,
                              target, description,
                              preserve_table_rows=(kind == 'table'))
                replaced += 1

            if progress:
                progress(index + 1, len(parts))

        document.save(docx_path)
    finally:
        pdf.close()
        import os
        for png in tmp:
            try:
                os.remove(png)
            except OSError:
                pass

    return replaced, missed
