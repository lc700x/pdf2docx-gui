import tempfile
import unittest
from pathlib import Path

import fitz
from docx import Document

import check


class DocxContentTests(unittest.TestCase):
    def test_checker_reads_image_descriptions_in_document_order(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            image_path = temp / 'pixel.png'
            pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 1, 1), False)
            pixmap.save(image_path)

            document = Document()
            document.add_paragraph('before')
            picture = document.add_paragraph().add_run().add_picture(
                str(image_path))
            picture._inline.docPr.set(
                'descr', 'cropped table alpha beta')
            document.add_paragraph('after')
            docx_path = temp / 'content.docx'
            document.save(docx_path)

            self.assertEqual(
                check.docx_words(docx_path),
                ['before', 'cropped', 'table', 'alpha', 'beta', 'after'])


if __name__ == '__main__':
    unittest.main()
