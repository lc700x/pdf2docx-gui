import unittest

from docx import Document

import snip_regions


class RegionAssignmentTests(unittest.TestCase):
    def test_paragraph_fragment_does_not_match_inside_a_crop_word(self):
        document = Document()
        paragraph = document.add_paragraph('fication')

        self.assertEqual(
            snip_regions._assign([paragraph._p], ['classification accuracy']),
            [None])

    def test_cropping_a_table_keeps_a_trailing_heading_row(self):
        document = Document()
        table = document.add_table(rows=0, cols=3)
        cropped = table.add_row().cells
        cropped[0].text = 'WTSP/fee-bee song'
        cropped[1].text = 'WTSP'
        cropped[2].text = '18/23'
        heading = table.add_row().cells
        heading[0].text = 'B'
        heading[1].text = 'Full results after the table'

        removed = snip_regions._remove_cropped_table_rows(
            table._tbl, 'WTSP/fee-bee song WTSP')

        self.assertEqual(removed, 1)
        self.assertEqual(len(table.rows), 1)
        self.assertEqual(table.cell(0, 1).text, 'Full results after the table')


if __name__ == '__main__':
    unittest.main()
