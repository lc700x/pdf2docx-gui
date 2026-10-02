import os
import struct
import tempfile
import unittest
import zipfile
import zlib
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Mm
from docx.text.paragraph import Paragraph

from latex_to_docx import (
    _create_reference_docx,
    _audit_converted_document,
    _constrain_nested_table_widths,
    _figure_reference_numbers,
    _keep_table_captions_with_tables,
    _normalize_minipage_tables,
    _normalize_table_captions,
    _normalize_table_markup,
    _normalize_subfigure_tabularx,
    _set_cell_padding,
    _set_table_column_widths,
    _preserve_latex_table_widths,
    _preserve_wide_subfigure_tables,
    _restore_figure_caption_labels,
    _restore_table_captions_and_missing_tables,
    _table_float_records,
    _wide_subfigure_linewidth_images,
    convert_latex_to_docx,
)


def _png_chunk(kind, data):
    body = kind + data
    return (struct.pack(">I", len(data)) + body
            + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF))


def _write_test_png(path):
    width, height = 300, 200
    header = struct.pack(">2I5B", width, height, 8, 6, 0, 0, 0)
    row = b"\x00" + b"\x33\x99\xff\xff" * width
    pixels = zlib.compress(row * height)
    with open(path, "wb") as image_file:
        image_file.write(b"\x89PNG\r\n\x1a\n")
        image_file.write(_png_chunk(b"IHDR", header))
        image_file.write(_png_chunk(b"IDAT", pixels))
        image_file.write(_png_chunk(b"IEND", b""))


class LatexConversionTests(unittest.TestCase):
    def test_shared_table_normalization_preserves_merged_and_multiline_cells(self):
        source = r"""\begin{tabularx}{\linewidth}{
>{\raggedright\arraybackslash}X
>{\centering\arraybackslash}p{0.2\linewidth}}
\toprule
\multicolumn{2}{@{}>{\centering\arraybackslash}p{\linewidth}}{\textbf{Accuracy}} \\
\makecell[b]{\textbf{Mean}\\(kHz)} & $\mu = 2$ \\
\cmidrule(lr){1-2}
Species & 7 \\
\midrule
\end{tabularx}"""

        normalized = _normalize_table_markup(source)

        self.assertIn(r"\begin{tabular}{lc}", normalized)
        self.assertIn(r"\multicolumn{2}{c}{\textbf{Accuracy}}", normalized)
        self.assertIn(
            r"\begin{tabular}{c}\textbf{Mean}\\(kHz)\end{tabular}",
            normalized)
        self.assertIn(r"$\mu = 2$", normalized)
        self.assertIn(r"\toprule", normalized)
        self.assertIn(r"\midrule", normalized)
        self.assertNotIn(r"\cmidrule", normalized)
        self.assertNotIn(r"\begin{tabularx}", normalized)

    def test_caption_normalization_trims_comments_before_adding_table_number(self):
        source = r"""\begin{table}
\caption{
% A source comment should not create a second label line.
   Classification accuracy across methods.
   Notes: values are percentages.
}
\begin{tabular}{lc}Method & Value \\\end{tabular}
\end{table}"""

        normalized = _normalize_table_captions(source)

        self.assertEqual(normalized.count("Classification accuracy"), 1)
        self.assertIn(
            r"\textbf{Table 1.} Classification accuracy across methods.",
            normalized)
        self.assertIn("Notes: values are percentages.", normalized)

    def test_nested_table_width_is_limited_by_merged_parent_cell(self):
        document = Document()
        parent = document.add_table(rows=1, cols=2)
        _set_table_column_widths(parent._tbl, [1.0, 1.0])
        merged = parent.cell(0, 0).merge(parent.cell(0, 1))
        _set_cell_padding(merged._tc, 40)
        nested = merged.add_table(rows=1, cols=2)
        _set_table_column_widths(nested._tbl, [2.0, 1.0])

        self.assertTrue(_constrain_nested_table_widths(parent._tbl))

        widths = [int(column.get(qn("w:w")))
                  for column in nested._tbl.find(qn("w:tblGrid"))]
        self.assertLessEqual(sum(widths), 2800)
        self.assertAlmostEqual(widths[0] / widths[1], 2, delta=0.01)

    def test_table_caption_and_notes_keep_with_the_table(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = os.path.join(directory, "caption.docx")
            document = Document()
            document.styles.add_style("Table Caption", 1)
            caption = document.add_paragraph("Table 1. Results")
            caption.style = "Table Caption"
            note = document.add_paragraph("Note: values are means.")
            document.add_table(rows=2, cols=2)
            document.save(output_path)

            _keep_table_captions_with_tables(output_path)

            converted = Document(output_path)
            caption_properties = converted.paragraphs[0]._p.find(qn("w:pPr"))
            note_properties = converted.paragraphs[1]._p.find(qn("w:pPr"))
            self.assertIsNotNone(caption_properties.find(qn("w:keepNext")))
            self.assertIsNotNone(note_properties.find(qn("w:keepNext")))

    def test_complex_table_converts_headers_values_math_and_formatting(self):
        source = r"""\documentclass{article}
\usepackage{makecell}
\begin{document}
\begin{table}
\caption{Complex table}
\begin{tabularx}{\textwidth}{>{\raggedright\arraybackslash}X >{\centering\arraybackslash}p{0.25\textwidth}}
\toprule
\multicolumn{2}{@{}>{\centering\arraybackslash}p{\textwidth}}{\textbf{Merged heading}} \\
\makecell[b]{\textbf{Mean}\\(kHz)} & $\mu = 2$ \\
\midrule
Speed & 7.5 \\
\bottomrule
\end{tabularx}
\end{table}
\end{document}"""
        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "complex.tex")
            output_path = os.path.join(directory, "complex.docx")
            with open(source_path, "w", encoding="utf-8") as source_file:
                source_file.write(source)

            convert_latex_to_docx(source_path, output_path)

            document = Document(output_path)
            table = next(table for table in document.tables
                         if table._tbl.find(qn("w:tblPr")).find(qn("w:tblStyle"))
                         is not None
                         and table._tbl.find(qn("w:tblPr")).find(qn("w:tblStyle"))
                         .get(qn("w:val")) == "Table")
            full_text = " ".join(node.text or ""
                                  for node in table._tbl.iter(qn("w:t")))
            self.assertIn("Merged heading", full_text)
            self.assertIn("Mean", full_text)
            self.assertIn("kHz", full_text)
            self.assertIn("Speed", full_text)
            self.assertIn("7.5", full_text)
            self.assertIsNotNone(table._tbl.find(".//" + qn("m:oMath")))
            heading_run = next(
                run for run in table._tbl.iter(qn("w:r"))
                if "Merged heading" in "".join(
                    node.text or "" for node in run.iter(qn("w:t"))))
            self.assertIsNotNone(heading_run.find(qn("w:rPr")))
            self.assertIsNotNone(heading_run.find(qn("w:rPr")).find(qn("w:b")))
            self.assertEqual(heading_run.find(qn("w:rPr")).find(qn("w:sz")).get(qn("w:val")), "20")
            header_rows = [row for row in table._tbl.findall(qn("w:tr"))
                           if row.find(qn("w:trPr")) is not None
                           and row.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None]
            self.assertTrue(header_rows)

    def test_wide_subfigure_linewidth_is_recognized(self):
        source = r"""\begin{subfigure}[t]{0.96\linewidth}
\includegraphics[width=\linewidth]{wide-panel.png}
\end{subfigure}"""

        image_count, widths = _wide_subfigure_linewidth_images(source)

        self.assertEqual(image_count, 1)
        self.assertEqual(widths, [(0, 0.96)])

    def test_long_citation_key_is_audited_by_bibliography_content(self):
        long_key = "heinrichAudioProtoPNetInterpretableDeep2024"
        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "references.bib")
            source_text = (
                r"\citep{" + long_key + r"}\bibliography{references}")
            with open(source_path, "w", encoding="utf-8") as bibliography:
                bibliography.write(
                    f"@article{{{long_key}, author={{A. Heinrich}}, "
                    "title={An Interpretable Deep Learning Model}, year={2024}}\n")
            output_path = os.path.join(directory, "references.docx")
            document = Document()
            document.styles.add_style("Bibliography", 1)
            document.add_paragraph("An Interpretable Deep Learning Model", "Bibliography")
            document.save(output_path)

            warnings = _audit_converted_document(
                source_text, directory, output_path)

        self.assertEqual(warnings, [])

    def test_missing_resources_are_reported_and_partial_docx_is_returned(self):
        source = r"""\documentclass{article}
\usepackage{graphicx}
\usepackage{natbib}
\begin{document}
See \citep{missing2024}.
\includegraphics[width=\linewidth]{figures/missing.png}
\bibliography{missing-references}
\end{document}"""
        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "main.tex")
            destination_path = os.path.join(directory, "main.docx")
            with open(source_path, "w", encoding="utf-8") as source_file:
                source_file.write(source)
            warnings = []
            progress = []

            result = convert_latex_to_docx(
                source_path, destination_path, on_warning=warnings.append,
                on_progress=progress.append)

            self.assertEqual(result, destination_path)
            self.assertTrue(os.path.isfile(destination_path))
            self.assertTrue(any("figures/missing.png" in warning
                                for warning in warnings))
            self.assertTrue(any("missing-references.bib" in warning
                                for warning in warnings))
            self.assertTrue(any("missing2024" in warning for warning in warnings))
            self.assertEqual(progress[0], 0)
            self.assertEqual(progress[-1], 100)
            self.assertEqual(progress, sorted(progress))

    def test_fatal_pandoc_failure_does_not_report_complete_progress(self):
        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "main.tex")
            output_path = os.path.join(directory, "main.docx")
            with open(source_path, "w", encoding="utf-8") as source_file:
                source_file.write(r"\documentclass{article}\begin{document}Text\end{document}")
            progress = []
            with patch("latex_to_docx.pypandoc.convert_file", side_effect=RuntimeError("Pandoc failed")):
                with self.assertRaisesRegex(RuntimeError, "Pandoc failed"):
                    convert_latex_to_docx(source_path, output_path, on_progress=progress.append)

            self.assertEqual(progress[0], 0)
            self.assertGreater(progress[-1], 0)
            self.assertLess(progress[-1], 100)
            self.assertFalse(os.path.exists(output_path))

    def test_subfigure_labels_reset_per_figure_and_preserve_caption_formatting(self):
        source = r"""\documentclass{article}
\begin{document}
\begin{figure}
\begin{subfigure}{0.5\textwidth}
\begin{tabular}{lc}Left & 1 \\ \end{tabular}
\caption[Left]{\textbf{Left panel} with $x=1$.}
\end{subfigure}
\begin{subfigure}{0.5\textwidth}
\begin{tabular}{lc}Right & 2 \\ \end{tabular}
\caption{Right panel.}
\end{subfigure}
\caption{First figure.}
\end{figure}
\begin{figure}
\begin{subfigure}{0.5\textwidth}
\begin{tabular}{lc}Upper & 3 \\ \end{tabular}
\caption{Upper panel.}
\end{subfigure}
\begin{subfigure}{0.5\textwidth}
\begin{tabular}{lc}Lower & 4 \\ \end{tabular}
\caption{Lower panel.}
\end{subfigure}
\caption{Second figure.}
\end{figure}
\end{document}"""
        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "panels.tex")
            output_path = os.path.join(directory, "panels.docx")
            with open(source_path, "w", encoding="utf-8") as source_file:
                source_file.write(source)
            convert_latex_to_docx(source_path, output_path)
            document = Document(output_path)
            captions = [Paragraph(element, document)
                        for table in document.tables
                        for element in table._tbl.iter(qn("w:p"))
                        if Paragraph(element, document).style.name == "Image Caption"]
            texts = [paragraph.text for paragraph in captions]
            self.assertEqual([text[:3] for text in texts], ["(a)", "(b)", "(a)", "(b)"])
            self.assertIn("Left panel", texts[0])
            self.assertTrue(any(run.bold and "Left panel" in run.text for run in captions[0].runs))
            self.assertIsNotNone(captions[0]._p.find(".//" + qn("m:oMath")))
            self.assertEqual([p.text for p in document.paragraphs if p.style.name == "Image Caption"],
                             ["Figure 1. First figure.", "Figure 2. Second figure."])

            _restore_figure_caption_labels(source, output_path)
            reopened = Document(output_path)
            repeated_texts = [Paragraph(element, reopened).text
                              for table in reopened.tables
                              for element in table._tbl.iter(qn("w:p"))
                              if Paragraph(element, reopened).style.name == "Image Caption"]
            self.assertEqual(repeated_texts, texts)

    def test_subfigure_tabularx_keeps_split_cells_and_declared_widths(self):
        source = r"""\documentclass{article}
\begin{document}
\begin{figure}
\begin{subfigure}{0.96\textwidth}
\begin{tabularx}{\linewidth}{>{\raggedright\arraybackslash}X >{\centering\arraybackslash}p{0.2\linewidth}}
Duration (s)
& 0.059 \\
\end{tabularx}
\caption{Measurements}
\end{subfigure}
\begin{subfigure}{\textwidth}
\begin{tabularx}{\linewidth}{>{\raggedright\arraybackslash}X >{\centering\arraybackslash}p{0.2\linewidth}}
Count
& 7 \\
\end{tabularx}
\caption{Counts}
\end{subfigure}
\caption{Measured values}
\end{figure}
\end{document}"""
        normalized = _normalize_subfigure_tabularx(source)
        self.assertIn("Duration (s) & 0.059", normalized)
        self.assertNotIn(r"\begin{tabularx}", normalized)

        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "main.tex")
            output_path = os.path.join(directory, "output.docx")
            with open(source_path, "w", encoding="utf-8") as source_file:
                source_file.write(source)
            convert_latex_to_docx(source_path, output_path)

            document = Document(output_path)
            figure_table = next(
                table for table in document.tables
                if table._tbl.find(qn("w:tblPr")).find(qn("w:tblStyle"))
                is not None
                and table._tbl.find(qn("w:tblPr")).find(qn("w:tblStyle"))
                .get(qn("w:val")) == "FigureTable")
            self.assertEqual((len(figure_table.rows), len(figure_table.columns)), (2, 1))

            def descendants(cell):
                for nested in cell.tables:
                    yield nested
                    for row in nested.rows:
                        for child in row.cells:
                            yield from descendants(child)

            nested_tables = [table for row in figure_table.rows
                             for table in descendants(row.cells[0])]
            data_table = next(
                table for table in nested_tables
                if len(table.columns) == 2
                and "0.059" in " ".join(
                    cell.text for row in table.rows for cell in row.cells))
            widths = [int(column.get(qn("w:w"))) / 1440
                      for column in data_table._tbl.find(qn("w:tblGrid"))]
            section = document.sections[0]
            text_width = (section.page_width - section.left_margin
                          - section.right_margin) / 914400
            self.assertAlmostEqual(widths[0], text_width * 0.96 * 0.8, delta=0.01)
            self.assertAlmostEqual(widths[1], text_width * 0.96 * 0.2, delta=0.01)

    def test_wide_subfigure_widths_ignore_commented_graphics(self):
        source = r"""\begin{subfigure}[t]{0.32\textwidth}
\includegraphics[width=\linewidth]{small.png}
\end{subfigure}
% \begin{subfigure}[t]{0.96\textwidth}
% \includegraphics[width=\linewidth]{commented.png}
% \end{subfigure}
\begin{subfigure}[t]{0.96\textwidth}
\includegraphics[width=\linewidth]{wide.png}
\end{subfigure}"""
        image_count, widths = _wide_subfigure_linewidth_images(source)
        self.assertEqual(image_count, 2)
        self.assertEqual(widths, [(1, 0.96)])

    def test_minipage_normalization_ignores_commented_environments(self):
        commented = r"""\begin{table}
% \begin{minipage}{\linewidth}
% \begin{tabularx}{\textwidth}{X}
% \end{tabularx}
% \end{minipage}
\begin{tabular}{l}alpha \\
\end{tabular}
\end{table}"""
        self.assertEqual(_normalize_minipage_tables(commented), commented)

        active = r"""\begin{table}
\begin{minipage}{\linewidth}
\begin{tabularx}{\textwidth}{X}
alpha \\
\end{tabularx}
\end{minipage}
\end{table}"""
        normalized = _normalize_minipage_tables(active)
        self.assertIn(r"\begin{tabular}{l}", normalized)
        self.assertNotIn(r"\begin{minipage}", normalized)
        self.assertNotIn(r"\begin{tabularx}", normalized)

    def test_recovers_native_table_when_pandoc_kept_only_its_caption(self):
        source = r"""\documentclass{article}
\begin{document}
\begin{table}
\caption{Unique table content}
\begin{tabular}{ll}
Header Alpha & Header Beta \\
Rare Zebra Marker & 42 \\
\end{tabular}
\begin{tabular}{ll}
Second Header Alpha & Second Header Beta \\
Rare Otter Marker & 17 \\
\end{tabular}
\end{table}
\end{document}"""
        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "main.tex")
            reference_path = os.path.join(directory, "reference.docx")
            output_path = os.path.join(directory, "output.docx")
            with open(source_path, "w", encoding="utf-8") as source_file:
                source_file.write(source)
            _create_reference_docx(source, reference_path)

            document = Document()
            caption = document.add_paragraph("Table 1. Unique table content")
            caption.style = "Caption"
            unrelated = document.add_table(rows=2, cols=2)
            unrelated.cell(0, 0).text = "Other heading"
            unrelated.cell(0, 1).text = "Different heading"
            unrelated.cell(1, 0).text = "Unrelated value"
            unrelated.cell(1, 1).text = "Another value"
            document.save(output_path)

            _restore_table_captions_and_missing_tables(
                source, _table_float_records(source), output_path, directory,
                reference_path, directory)

            converted = Document(output_path)
            self.assertEqual(len(converted.tables), 3)
            self.assertEqual(
                sum(p.text.startswith("Table 1.") for p in converted.paragraphs), 1)
            self.assertTrue(any(
                "Rare Zebra Marker" in cell.text
                for table in converted.tables
                for row in table.rows for cell in row.cells))
            self.assertTrue(any(
                "Rare Otter Marker" in cell.text
                for table in converted.tables
                for row in table.rows for cell in row.cells))

    def test_figure_references_use_figure_and_subfigure_numbers(self):
        source = r"""\begin{figure}
\begin{subfigure}{0.5\textwidth}
\includegraphics{panel.png}
\caption{Panel caption}\label{fig:panel}
\end{subfigure}
\caption{Overall caption}\label{fig:overall}
\end{figure}"""
        self.assertEqual(_figure_reference_numbers(source), {
            "fig:panel": "1a",
            "fig:overall": "1",
        })

    def test_tabularx_widths_preserve_fixed_and_stretch_columns(self):
        source = r"""\documentclass{article}
\begin{document}
\begin{table}
\caption{Species}
\begin{tabularx}{\linewidth}{>{\raggedright\arraybackslash}p{0.38\linewidth} X >{\centering\arraybackslash}p{0.18\linewidth}}
Common Name & Latin Name & Code \\
Alder Flycatcher & Empidonax alnorum & ALFL \\
\end{tabularx}
\end{table}
\end{document}"""
        with tempfile.TemporaryDirectory() as directory:
            output_path = os.path.join(directory, "widths.docx")
            document = Document()
            document.sections[0].page_width = Mm(210)
            document.sections[0].page_height = Mm(297)
            document.sections[0].left_margin = Inches(1)
            document.sections[0].right_margin = Inches(1)
            document.add_paragraph("Table 1. Species")
            document.add_table(rows=2, cols=3)
            document.save(output_path)

            _preserve_latex_table_widths(source, output_path)

            converted = Document(output_path)
            widths = [int(column.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"))
                      / 1440 for column in converted.tables[0]._tbl.tblGrid]
            text_width = 210 / 25.4 - 2
            self.assertAlmostEqual(sum(widths), text_width, delta=0.01)
            self.assertAlmostEqual(widths[0], text_width * 0.38, delta=0.01)
            self.assertAlmostEqual(widths[1], text_width * 0.44, delta=0.01)
            self.assertAlmostEqual(widths[2], text_width * 0.18, delta=0.01)

    def test_conversion_keeps_text_tables_math_figures_and_citations(self):
        with tempfile.TemporaryDirectory() as directory:
            image_dir = os.path.join(directory, "figures")
            os.mkdir(image_dir)
            image_path = os.path.join(image_dir, "pixel.png")
            _write_test_png(image_path)

            source_path = os.path.join(directory, "main.tex")
            with open(source_path, "w", encoding="utf-8") as source_file:
                source_file.write(r"""\documentclass[12pt,a4paper]{article}
\usepackage[margin=1in]{geometry}
\usepackage{graphicx}
\usepackage{subcaption}
\usepackage{natbib}
\begin{document}
\section{Editable section}
An editable paragraph with inline math $x^2$, citation \citep{smith2020}, and 50\% coverage.
\begin{equation}x = 2\end{equation}
\begin{table}
\caption{Numeric results}\label{tab:numeric}
\begin{tabular}{lc}
Name & Value \\
alpha & 3 \\
\end{tabular}
\end{table}
\begin{figure}
\includegraphics[
% TeX comments inside a graphics option must not break Pandoc parsing.
width=0.5\textwidth
]{figures/pixel.png}
\caption{Pixel plot}\label{fig:pixel}
\end{figure}
\begin{figure}
\begin{subfigure}[t]{0.96\textwidth}
\includegraphics[width=\linewidth]{figures/pixel.png}
\caption{Wide pixel plot}
\end{subfigure}
\end{figure}
\bibliographystyle{plainnat}
\bibliography{references}
See Table~\ref{tab:numeric} and Figure~\ref{fig:pixel}.\par
\appendix
\section{Appendix material}\label{app:details}
See Appendix~\ref{app:details}.
\end{document}
""")
            with open(os.path.join(directory, "references.bib"), "w",
                      encoding="utf-8") as bibliography:
                bibliography.write("""@article{smith2020,
  author = {Jane Smith},
  title = {A Useful Paper},
  journal = {Test Journal},
  year = {2020}
}
""")

            destination_path = os.path.join(directory, "main_from_tex.docx")
            statuses = []
            progress = []
            with open(source_path, "rb") as source_file:
                original_source = source_file.read()
            result = convert_latex_to_docx(
                source_path, destination_path, on_status=statuses.append,
                on_progress=progress.append)

            self.assertEqual(result, destination_path)
            with open(source_path, "rb") as source_file:
                self.assertEqual(source_file.read(), original_source)
            self.assertIn("Converting LaTeX with Pandoc...", statuses)
            self.assertIn("Checking DOCX content...", statuses)
            self.assertEqual(statuses[-1], "Finalizing DOCX...")
            self.assertEqual(progress, [round(stage * 100 / 14) for stage in range(15)])
            document = Document(destination_path)
            document_text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables
                   for row in table.rows for cell in row.cells])
            self.assertIn("Editable section", document_text)
            self.assertIn("editable paragraph", document_text.lower())
            self.assertIn("50% coverage", document_text)
            self.assertIn("A Useful Paper", document_text)
            self.assertIn("Pixel plot", document_text)
            self.assertIn("Wide pixel plot", document_text)
            self.assertIn("Table 1. Numeric results", document_text)
            normalized_text = " ".join(document_text.replace("\xa0", " ").split())
            self.assertIn("See Table 1 and Figure 1.", normalized_text)
            self.assertIn("See Appendix A.", normalized_text)
            appendix_heading = next(
                paragraph for paragraph in document.paragraphs
                if paragraph.style.name == "Heading 1"
                and "Appendix material" in paragraph.text)
            self.assertTrue(appendix_heading.text.startswith("Appendix A "))
            figure_captions = [
                paragraph for paragraph in document.paragraphs
                if paragraph.style.name == "Image Caption"]
            self.assertTrue(any(
                paragraph.text == "Figure 1. Pixel plot"
                for paragraph in figure_captions))
            self.assertTrue(any(
                paragraph.runs and paragraph.runs[0].text == "Figure 1."
                and paragraph.runs[0].bold
                for paragraph in figure_captions))
            table_caption = next(
                paragraph for paragraph in document.paragraphs
                if paragraph.text.startswith("Table 1."))
            self.assertEqual(table_caption.style.name, "Table Caption")
            heading = next(
                paragraph for paragraph in document.paragraphs
                if paragraph.text.endswith("Editable section"))
            self.assertEqual(heading.style.name, "Heading 1")
            self.assertEqual(len(document.tables), 2)
            self.assertTrue(any(
                any(cell.text == "alpha" for row in table.rows for cell in row.cells)
                for table in document.tables))
            self.assertEqual(len(document.inline_shapes), 2)
            self.assertGreater(document.inline_shapes[0].width, 0)
            self.assertGreater(document.inline_shapes[0].height, 0)
            expected_width = (210 / 25.4 - 2) / 2
            self.assertAlmostEqual(
                document.inline_shapes[0].width / 914400,
                expected_width, delta=0.1)
            self.assertAlmostEqual(
                document.inline_shapes[0].width
                / document.inline_shapes[0].height, 1.5, places=2)
            self.assertAlmostEqual(
                document.inline_shapes[1].width / 914400,
                expected_width * 2 * 0.96, delta=0.1)
            self.assertAlmostEqual(
                document.sections[0].page_width / 914400, 210 / 25.4, places=1)
            self.assertAlmostEqual(
                document.sections[0].left_margin / 914400, 1.0, places=2)
            self.assertEqual(document.styles["Normal"].font.name, "Times New Roman")
            self.assertEqual(document.styles["Normal"].font.size.pt, 12)
            self.assertEqual(
                document.styles["Image Caption"].font.size.pt, 10)
            self.assertEqual(
                document.styles["Heading 1"].paragraph_format.space_before.pt,
                12)
            with zipfile.ZipFile(destination_path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
                self.assertIn("<m:oMath", xml)

    def test_missing_source_reports_clear_error(self):
        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "missing.tex")
            destination_path = os.path.join(directory, "missing.docx")
            with self.assertRaisesRegex(FileNotFoundError, "source file not found"):
                convert_latex_to_docx(source_path, destination_path)


if __name__ == "__main__":
    unittest.main()
