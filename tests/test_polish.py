import tempfile
import unittest
from pathlib import Path

import fitz
from docx import Document

import polish


class SourceSpacingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.pdf_path = Path(self.temp.name) / 'source.pdf'
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text(
            (72, 72),
            'Benjamin M Bolker recent years; provide a compact setup set up.')
        pdf.save(self.pdf_path)
        pdf.close()

    def tearDown(self):
        self.temp.cleanup()

    def test_restores_rare_source_word_boundaries(self):
        source_words, source_pairs = polish.source_word_boundaries(
            self.pdf_path)
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Benjamin')
        paragraph.add_run('M')
        paragraph.add_run('Bolker')
        paragraph.add_run('recent')
        paragraph.add_run('years.')
        paragraph = document.add_paragraph()
        paragraph.add_run('provide')
        paragraph.add_run('a')
        paragraph.add_run('compact.')

        restored = polish.restore_spaces(
            document, set(), source_words, source_pairs)

        self.assertEqual(restored, 6)
        self.assertEqual(document.paragraphs[0].text,
                         'Benjamin M Bolker recent years.')
        self.assertEqual(document.paragraphs[1].text, 'provide a compact.')

    def test_does_not_split_a_valid_source_word_or_hyphenation(self):
        source_words, source_pairs = polish.source_word_boundaries(
            self.pdf_path)
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('set')
        paragraph.add_run('up')
        paragraph.add_run(' high-')
        paragraph.add_run('quality')

        polish.restore_spaces(document, set(), source_words, source_pairs)

        self.assertEqual(paragraph.text, 'setup high-quality')


if __name__ == '__main__':
    unittest.main()
